#!/usr/bin/env python3
"""Gera planilhas FICTICIAS para demonstrar o dashboard.

Todos os nomes, setores e valores sao inventados (seed fixa = resultado
reprodutivel). Nenhum dado real e usado.

Saida:
  bd_exemplo/Entradas/Entradas_AAAA.xlsx   (aba "Entradas")
      Data | Entradas do dia | Receita de Vendas imóveis | Banco
  bd_exemplo/Saídas/Saidas_AAAA.xlsx       (aba "AAAA")
      Data | Departamento | Modalidade | Natureza | Descriçao das Naturezas |
      Favorecido | Beneficiario | Banco | Valor

Uso:
  pip install openpyxl
  python gerar_dados_exemplo.py
"""
import calendar
import datetime as dt
import os
import random

from openpyxl import Workbook

SEED = 20260101
ANOS = [2023, 2024, 2025, 2026]
ULTIMO_MES_2026 = 9            # gera 2026 ate setembro
RAIZ = os.path.join(os.path.dirname(os.path.abspath(__file__)), "bd_exemplo")

# setor -> (peso no gasto total, crescimento anual)
SETORES = {
    "ADMINISTRATIVO": (0.16, 0.05),
    "OPERACOES": (0.15, 0.08),
    "TECNOLOGIA": (0.12, 0.10),
    "EVENTOS": (0.07, 0.04),
    "RH": (0.14, 0.05),
    "FINANCEIRO": (0.05, 0.03),
    "JURIDICO": (0.04, 0.02),
    "MARKETING": (0.06, 0.06),
    "COMPRAS": (0.05, 0.04),
    "LOGISTICA": (0.05, 0.05),
    "MANUTENCAO": (0.07, 0.07),
    "VIAGENS": (0.04, 0.03),
}

MODALIDADES = ["Transferência", "Boleto", "Débito automático", "Cartão corporativo"]
NATUREZAS = {
    3101: "Serviços de terceiros", 3102: "Material de consumo", 3103: "Manutenção predial",
    3104: "Licenças de software", 3105: "Passagens e hospedagem", 3106: "Folha e encargos",
    3107: "Eventos e locação", 3108: "Publicidade", 3109: "Assessoria técnica",
}
FORNECEDORES = [f"Fornecedor Exemplo {chr(65 + i)}" for i in range(20)]
BENEFICIARIOS = [f"Beneficiário Fictício {i:02d}" for i in range(1, 31)]
BANCOS = ["Banco Alfa", "Banco Beta", "Banco Gama"]


def dias_uteis(ano, mes):
    n = calendar.monthrange(ano, mes)[1]
    return [dt.datetime(ano, mes, d) for d in range(1, n + 1)
            if dt.date(ano, mes, d).weekday() < 5]


def meses_do_ano(ano):
    return range(1, (ULTIMO_MES_2026 if ano == 2026 else 12) + 1)


def sazonalidade(mes):
    # leve alta no fim do ano
    return {11: 1.08, 12: 1.15, 1: 0.95, 2: 0.93}.get(mes, 1.0)


def total_saidas_mes(ano, mes, rnd):
    base = 2_400_000 * (1.06 ** (ano - 2023))
    return base * sazonalidade(mes) * rnd.uniform(0.94, 1.06)


def gerar_entradas(ano, rnd):
    wb = Workbook()
    ws = wb.active
    ws.title = "Entradas"
    ws.append(["Data", "Entradas do dia", "Receita de Vendas imóveis", "Banco"])
    for mes in meses_do_ano(ano):
        alvo = total_saidas_mes(ano, mes, rnd) * rnd.uniform(1.02, 1.18)
        dias = dias_uteis(ano, mes)
        pesos = [rnd.uniform(0.4, 1.8) for _ in dias]
        soma = sum(pesos)
        for d, p in zip(dias, pesos):
            valor = round(alvo * p / soma, 2)
            imov = round(valor * rnd.uniform(0, 0.12), 2) if rnd.random() < 0.15 else 0
            ws.append([d, valor, imov, rnd.choice(BANCOS)])
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 26
    ws.column_dimensions["D"].width = 14
    for row in ws.iter_rows(min_row=2, max_col=1):
        row[0].number_format = "dd/mm/yyyy"
    pasta = os.path.join(RAIZ, "Entradas")
    os.makedirs(pasta, exist_ok=True)
    wb.save(os.path.join(pasta, f"Entradas_{ano}.xlsx"))


def gerar_saidas(ano, rnd):
    wb = Workbook()
    ws = wb.active
    ws.title = str(ano)
    ws.append(["Data", "Departamento", "Modalidade", "Natureza", "Descriçao das Naturezas",
               "Favorecido", "Beneficiario", "Banco", "Valor"])
    for mes in meses_do_ano(ano):
        total_mes = total_saidas_mes(ano, mes, rnd)
        dias = dias_uteis(ano, mes)
        # normaliza os pesos (ja com crescimento anual por setor)
        pesos = {s: w * ((1 + g) ** (ano - 2023)) * rnd.uniform(0.9, 1.1)
                 for s, (w, g) in SETORES.items()}
        soma = sum(pesos.values())
        for setor, peso in pesos.items():
            alvo = total_mes * peso / soma
            n = rnd.randint(14, 26)
            partes = [rnd.uniform(0.2, 2.2) for _ in range(n)]
            sp = sum(partes)
            acumulado = 0.0
            for i, p in enumerate(partes):
                valor = round(alvo * p / sp, 2) if i < n - 1 else round(alvo - acumulado, 2)
                acumulado += valor
                nat = rnd.choice(list(NATUREZAS))
                ws.append([rnd.choice(dias), setor, rnd.choice(MODALIDADES), nat, NATUREZAS[nat],
                           rnd.choice(FORNECEDORES), rnd.choice(BENEFICIARIOS),
                           rnd.choice(BANCOS), valor])
    for col, larg in zip("ABCDEFGHI", [12, 18, 20, 10, 26, 24, 24, 14, 14]):
        ws.column_dimensions[col].width = larg
    for row in ws.iter_rows(min_row=2, max_col=1):
        row[0].number_format = "dd/mm/yyyy"
    pasta = os.path.join(RAIZ, "Saídas")
    os.makedirs(pasta, exist_ok=True)
    wb.save(os.path.join(pasta, f"Saidas_{ano}.xlsx"))


def main():
    rnd = random.Random(SEED)
    for ano in ANOS:
        gerar_entradas(ano, rnd)
        gerar_saidas(ano, rnd)
        print(f"{ano}: ok")
    print("Planilhas fictícias geradas em:", RAIZ)


if __name__ == "__main__":
    main()
